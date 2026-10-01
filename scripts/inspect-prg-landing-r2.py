#!/usr/bin/env python3
"""Inspect the current GUGiK PRG Landing objects and XLSX schema from private R2.

Read-only. No R2 object, Iceberg table, or control state is modified.
"""

from __future__ import annotations

import json
import os
from pathlib import Path
import tempfile
from typing import Any

import boto3
from botocore.config import Config
from openpyxl import load_workbook

PREFIX = "01_landing/gugik_prg/"


def req(name: str) -> str:
    value=os.environ.get(name,"").strip()
    if not value: raise RuntimeError(f"missing {name}")
    return value


def client():
    return boto3.client(
        "s3",
        endpoint_url=req("R2_S3_ENDPOINT").rstrip("/"),
        region_name="auto",
        aws_access_key_id=req("CLOUDFLARE_R2_ACCESS_KEY_ID"),
        aws_secret_access_key=req("CLOUDFLARE_R2_SECRET_ACCESS_KEY"),
        config=Config(signature_version="s3v4",retries={"max_attempts":8,"mode":"standard"}),
    )


def clean(value: Any) -> Any:
    if value is None or isinstance(value,(str,int,float,bool)):
        return value
    return str(value)


def main() -> int:
    s3=client(); bucket=req("R2_LANDING_BUCKET")
    objects=[]
    for page in s3.get_paginator("list_objects_v2").paginate(Bucket=bucket,Prefix=PREFIX):
        for item in page.get("Contents",[]):
            key=str(item["Key"]); size=int(item.get("Size",0))
            if key.endswith("/") and size==0: continue
            objects.append({"key":key,"size":size})
    objects.sort(key=lambda x:x["key"])
    xlsx=[x for x in objects if x["key"].lower().endswith(".xlsx")]
    zips=[x for x in objects if x["key"].lower().endswith(".zip")]
    if len(xlsx)!=1 or len(zips)!=1:
        raise RuntimeError(f"expected one PRG ZIP and one XLSX; observed zip={len(zips)} xlsx={len(xlsx)}")

    with tempfile.TemporaryDirectory(prefix="zohelo-prg-inspect-") as td:
        path=Path(td)/Path(xlsx[0]["key"]).name
        s3.download_file(bucket,xlsx[0]["key"],str(path))
        if path.stat().st_size!=xlsx[0]["size"]:
            raise RuntimeError("PRG XLSX size changed while downloading")
        book=load_workbook(path,read_only=True,data_only=True)
        sheets=[]
        try:
            for ws in book.worksheets:
                rows=[]
                for idx,row in enumerate(ws.iter_rows(values_only=True),start=1):
                    rows.append([clean(value) for value in row])
                    if idx>=12: break
                sheets.append({
                    "title":ws.title,
                    "max_row":ws.max_row,
                    "max_column":ws.max_column,
                    "sample_rows":rows,
                })
        finally:
            book.close()

    print(json.dumps({
        "result":"pass",
        "source":"gugik_prg",
        "landing_objects":objects,
        "xlsx":xlsx[0],
        "zip":zips[0],
        "workbook_sheets":sheets,
        "writes_performed":False,
    },ensure_ascii=False,sort_keys=True))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
