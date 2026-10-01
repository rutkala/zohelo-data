#!/usr/bin/env python3
"""Promote the missing PRG surface-area XLSX from Landing into Iceberg.

This source-specific lifecycle step:
1. reads the exact current PRG XLSX from private R2 Landing;
2. hashes the native XLSX and parses its single worksheet deterministically;
3. writes one source-shaped Bronze Parquet with explicit native lineage;
4. commits bronze.prg_surface_areas and zero-copy Silver/Gold tables;
5. validates representative R2 SQL reads;
6. archives the native XLSX with exact SHA-256 readback;
7. deletes only that XLSX from Landing when no source-control reference blocks it.

The PRG boundary ZIP remains untouched in Landing pending its separate lineage proof.
"""

from __future__ import annotations

from datetime import date, datetime, timezone
from hashlib import sha256
import json
import os
from pathlib import Path
import tempfile
from typing import Any
from urllib.parse import urlsplit

import boto3
from botocore.config import Config
from botocore.exceptions import ClientError
from openpyxl import load_workbook
import pyarrow as pa
import pyarrow.parquet as pq
from pyiceberg.catalog.rest import RestCatalog
from pyiceberg.exceptions import NamespaceAlreadyExistsError, NoSuchTableError, TableAlreadyExistsError
from pyiceberg.manifest import DataFile, DataFileContent, FileFormat
from pyiceberg.table import TableProperties
from pyiceberg.typedef import Record
import requests

SOURCE="gugik_prg"
LANDING_PREFIX="01_landing/gugik_prg/"
BRONZE_ID=("bronze","prg_surface_areas")
SILVER_ID=("silver","prg_surface_areas")
GOLD_ID=("gold","fact_prg_surface_areas")
SNAPSHOT_DATE=date(2026,1,1)


def req(name: str) -> str:
    value=os.environ.get(name,"").strip()
    if not value: raise RuntimeError(f"missing {name}")
    return value


def s3_client():
    return boto3.client(
        "s3",endpoint_url=req("R2_S3_ENDPOINT").rstrip("/"),region_name="auto",
        aws_access_key_id=req("CLOUDFLARE_R2_ACCESS_KEY_ID"),
        aws_secret_access_key=req("CLOUDFLARE_R2_SECRET_ACCESS_KEY"),
        config=Config(signature_version="s3v4",retries={"max_attempts":8,"mode":"standard"}))


def catalog():
    return RestCatalog(
        name="zohelo",warehouse=req("R2_WAREHOUSE"),uri=req("R2_CATALOG_URI"),
        token=req("R2_DATA_CATALOG_TOKEN"),
        **{"s3.access-key-id":req("CLOUDFLARE_R2_ACCESS_KEY_ID"),
           "s3.secret-access-key":req("CLOUDFLARE_R2_SECRET_ACCESS_KEY"),
           "s3.endpoint":req("R2_S3_ENDPOINT"),"s3.region":"auto"})


def hash_file(path: Path) -> str:
    digest=sha256()
    with path.open("rb") as handle:
        for block in iter(lambda:handle.read(1024*1024),b""):
            digest.update(block)
    return digest.hexdigest()


def hash_object(s3,bucket: str,key: str) -> tuple[str,int]:
    response=s3.get_object(Bucket=bucket,Key=key)
    body=response["Body"]; digest=sha256(); size=0
    try:
        while True:
            chunk=body.read(1024*1024)
            if not chunk: break
            digest.update(chunk); size+=len(chunk)
    finally:
        body.close()
    return digest.hexdigest(),size


def list_xlsx(s3,bucket: str) -> dict[str,Any]:
    items=[]
    for page in s3.get_paginator("list_objects_v2").paginate(Bucket=bucket,Prefix=LANDING_PREFIX):
        for obj in page.get("Contents",[]):
            key=str(obj["Key"]); size=int(obj.get("Size",0))
            if key.lower().endswith(".xlsx") and size>0:
                items.append({"key":key,"size":size})
    if len(items)!=1:
        raise RuntimeError(f"expected exactly one PRG XLSX, observed {len(items)}")
    return items[0]


def parse_workbook(path: Path, source_key: str, source_sha: str) -> pa.Table:
    book=load_workbook(path,read_only=True,data_only=True)
    try:
        if book.sheetnames!=["Arkusz1"]:
            raise RuntimeError(f"unexpected PRG workbook sheets: {book.sheetnames}")
        ws=book["Arkusz1"]
        rows=ws.iter_rows(values_only=True)
        header=next(rows,None)
        expected=("TERYT","Nazwa jednostki","Powierzchnia [m2]","Powierzchnia [ha]","Powierzchnia [km2]")
        observed=tuple(header[:5]) if header else ()
        if observed!=expected:
            raise RuntimeError(f"unexpected PRG XLSX header: {observed}")

        values=[]
        for row_number,row in enumerate(rows,start=2):
            if not row or all(value is None for value in row[:5]):
                continue
            raw_code=row[0]
            name=row[1]
            m2=row[2]; ha=row[3]; km2=row[4]
            if raw_code is None or name is None or any(v is None for v in (m2,ha,km2)):
                raise RuntimeError(f"incomplete PRG XLSX row {row_number}")
            code_source=str(raw_code).rstrip()
            code="".join(code_source.split())
            if not code or not code.isdigit():
                raise RuntimeError(f"invalid TERYT code at PRG XLSX row {row_number}")
            try:
                area_m2=int(m2); area_ha=int(ha); area_km2=int(km2)
            except (TypeError,ValueError) as exc:
                raise RuntimeError(f"invalid area measure at PRG XLSX row {row_number}") from exc
            values.append({
                "teryt_code_source":code_source,
                "teryt_code":code,
                "unit_name":str(name),
                "area_m2":area_m2,
                "area_ha":area_ha,
                "area_km2":area_km2,
                "snapshot_date":SNAPSHOT_DATE,
                "source_sheet":"Arkusz1",
                "source_row_number":row_number,
                "source_object_key":source_key,
                "source_file_sha256":source_sha,
            })
    finally:
        book.close()

    if len(values)!=4351:
        raise RuntimeError(f"expected 4351 PRG surface-area rows, observed {len(values)}")

    schema=pa.schema([
        ("teryt_code_source",pa.string()),
        ("teryt_code",pa.string()),
        ("unit_name",pa.string()),
        ("area_m2",pa.int64()),
        ("area_ha",pa.int64()),
        ("area_km2",pa.int64()),
        ("snapshot_date",pa.date32()),
        ("source_sheet",pa.string()),
        ("source_row_number",pa.int64()),
        ("source_object_key",pa.string()),
        ("source_file_sha256",pa.string()),
    ])
    return pa.Table.from_pylist(values,schema=schema)


def ensure_namespace(cat,namespace: str) -> None:
    try:
        cat.create_namespace(namespace)
    except NamespaceAlreadyExistsError:
        pass


def minimal_data_file(table,uri: str,rows: int,size: int) -> DataFile:
    return DataFile.from_args(
        _table_format_version=int(table.metadata.format_version),
        content=DataFileContent.DATA,file_path=uri,file_format=FileFormat.PARQUET,
        partition=Record(),record_count=rows,file_size_in_bytes=size,
        spec_id=int(table.metadata.default_spec_id))


def ensure_table(cat,identifier: tuple[str,str],schema: pa.Schema,uri: str,rows: int,size: int,properties: dict[str,str]):
    ensure_namespace(cat,identifier[0])
    try:
        table=cat.load_table(identifier)
    except NoSuchTableError:
        try:
            table=cat.create_table(identifier,schema=schema,properties=properties)
        except TableAlreadyExistsError:
            table=cat.load_table(identifier)

    existing=list(table.scan().plan_files())
    paths={str(task.file.file_path) for task in existing}
    if paths and paths!={uri}:
        raise RuntimeError(f"unexpected existing membership in {'.'.join(identifier)}: {paths}")
    if not paths:
        with table.transaction() as tx:
            if tx.table_metadata.name_mapping() is None:
                tx.set_properties(**{TableProperties.DEFAULT_NAME_MAPPING:
                    tx.table_metadata.schema().name_mapping.model_dump_json()})
            with tx.update_snapshot(snapshot_properties={
                "zohelo.operation":"prg-surface-area-landing-to-iceberg",
                "zohelo.source":SOURCE,
            }).fast_append() as append:
                append.append_data_file(minimal_data_file(table,uri,rows,size))
        table=cat.load_table(identifier)

    files=list(table.scan().plan_files())
    if len(files)!=1:
        raise RuntimeError(f"unexpected data-file count in {'.'.join(identifier)}")
    task=files[0]
    if (
        str(task.file.file_path)!=uri
        or int(task.file.record_count)!=rows
        or int(task.file.file_size_in_bytes)!=size
    ):
        raise RuntimeError(f"membership reconciliation failed for {'.'.join(identifier)}")

    desired=dict(properties)
    desired["zohelo.source.file-sha256"]=properties["zohelo.source.file-sha256"]
    if any(table.properties.get(k)!=v for k,v in desired.items()):
        with table.transaction() as tx:
            tx.set_properties(**desired)
        table=cat.load_table(identifier)
    return table


def r2_sql(query: str) -> None:
    url=(f"https://api.sql.cloudflarestorage.com/api/v1/accounts/"
         f"{req('CLOUDFLARE_ACCOUNT_ID')}/r2-sql/query/{req('R2_LAKEHOUSE_BUCKET')}")
    response=requests.post(
        url,
        headers={"Authorization":f"Bearer {req('R2_DATA_CATALOG_TOKEN')}",
                 "Content-Type":"application/json"},
        json={"query":query},timeout=120)
    if response.status_code!=200:
        raise RuntimeError(f"R2 SQL {response.status_code}: {response.text[:400]}")
    payload=response.json()
    if not payload.get("success",False):
        raise RuntimeError(f"R2 SQL failed: {payload}")


def control_references(s3,lakehouse_bucket: str,landing_bucket: str,key: str) -> list[str]:
    needles=(key.encode("utf-8"),f"r2://{landing_bucket}/{key}".encode("utf-8"))
    matches=[]
    prefix="06_control/source_campaigns/gugik_prg/"
    for page in s3.get_paginator("list_objects_v2").paginate(Bucket=lakehouse_bucket,Prefix=prefix):
        for obj in page.get("Contents",[]):
            control_key=str(obj["Key"]); size=int(obj.get("Size",0))
            if size<=0 or size>8*1024*1024: continue
            body=s3.get_object(Bucket=lakehouse_bucket,Key=control_key)["Body"].read()
            if any(needle in body for needle in needles):
                matches.append(control_key)
    return sorted(matches)


def put_immutable_json(s3,bucket: str,key: str,value: dict[str,Any]) -> None:
    raw=(json.dumps(value,indent=2,sort_keys=True)+"\n").encode("utf-8")
    try:
        existing=s3.get_object(Bucket=bucket,Key=key)["Body"].read()
    except Exception:
        existing=None
    if existing is not None:
        if existing!=raw: raise RuntimeError(f"immutable receipt conflict: {key}")
        return
    s3.put_object(Bucket=bucket,Key=key,Body=raw,ContentType="application/json",
                  Metadata={"sha256":sha256(raw).hexdigest()})


def object_exists(s3,bucket: str,key: str) -> bool:
    try:
        s3.head_object(Bucket=bucket,Key=key); return True
    except ClientError as exc:
        status=int(exc.response.get("ResponseMetadata",{}).get("HTTPStatusCode",0) or 0)
        if status==404: return False
        raise


def main() -> int:
    s3=s3_client(); cat=catalog()
    landing_bucket=req("R2_LANDING_BUCKET"); lakehouse_bucket=req("R2_LAKEHOUSE_BUCKET")
    item=list_xlsx(s3,landing_bucket)

    with tempfile.TemporaryDirectory(prefix="zohelo-prg-surface-") as td:
        root=Path(td)
        xlsx=root/"wykaz_powierzchni_2026.xlsx"
        s3.download_file(landing_bucket,item["key"],str(xlsx))
        if xlsx.stat().st_size!=item["size"]:
            raise RuntimeError("PRG XLSX size changed while downloading")
        source_sha=hash_file(xlsx)
        table=parse_workbook(xlsx,item["key"],source_sha)
        parquet=root/"prg_surface_areas.parquet"
        pq.write_table(table,parquet,compression="zstd")
        rows=table.num_rows; parquet_size=parquet.stat().st_size

        data_key=f"02_bronze/prg_surface_areas/data/{source_sha}.parquet"
        try:
            head=s3.head_object(Bucket=lakehouse_bucket,Key=data_key)
            present=int(head.get("ContentLength",-1))==parquet_size
        except Exception:
            present=False
        if not present:
            s3.upload_file(str(parquet),lakehouse_bucket,data_key,
                           ExtraArgs={"ContentType":"application/vnd.apache.parquet",
                                      "Metadata":{"source-sha256":source_sha}})
        head=s3.head_object(Bucket=lakehouse_bucket,Key=data_key)
        if int(head.get("ContentLength",-1))!=parquet_size:
            raise RuntimeError("PRG Bronze Parquet upload size mismatch")

    uri=f"s3://{lakehouse_bucket}/{data_key}"
    common={
        "zohelo.source":SOURCE,
        "zohelo.source.object-key":item["key"],
        "zohelo.source.file-sha256":source_sha,
        "zohelo.source.snapshot-date":SNAPSHOT_DATE.isoformat(),
        "zohelo.ready":"partial-current-data",
    }
    bronze=ensure_table(cat,BRONZE_ID,table.schema,uri,rows,parquet_size,{
        **common,"zohelo.transform.mode":"landing-to-bronze-v1"})
    silver=ensure_table(cat,SILVER_ID,table.schema,uri,rows,parquet_size,{
        **common,"zohelo.transform.mode":"zero-copy-promotion-v1",
        "zohelo.source.table":"bronze.prg_surface_areas"})
    gold=ensure_table(cat,GOLD_ID,table.schema,uri,rows,parquet_size,{
        **common,"zohelo.transform.mode":"zero-copy-promotion-v1",
        "zohelo.source.table":"silver.prg_surface_areas"})

    for query in (
        "SELECT * FROM bronze.prg_surface_areas LIMIT 1;",
        "SELECT * FROM silver.prg_surface_areas LIMIT 1;",
        "SELECT * FROM gold.fact_prg_surface_areas LIMIT 1;",
    ):
        r2_sql(query)

    archive_key=f"05_archive/gugik_prg/{source_sha}/{Path(item['key']).name}"
    if not object_exists(s3,landing_bucket,archive_key):
        s3.copy_object(Bucket=landing_bucket,Key=archive_key,
                       CopySource={"Bucket":landing_bucket,"Key":item["key"]})
    archive_sha,archive_size=hash_object(s3,landing_bucket,archive_key)
    if archive_sha!=source_sha or archive_size!=item["size"]:
        raise RuntimeError("PRG XLSX Archive readback mismatch")

    blockers=control_references(s3,lakehouse_bucket,landing_bucket,item["key"])
    receipt_key=f"06_control/lineage/gugik_prg/surface-area-{source_sha}.json"
    receipt_uri=f"s3://{lakehouse_bucket}/{receipt_key}"
    receipt={
        "format_version":1,
        "kind":"landing_bronze_lineage_receipt",
        "source_id":SOURCE,
        "dataset_id":"prg_surface_areas",
        "input_set_sha256":source_sha,
        "landing_inputs":[{"key":item["key"],"size":item["size"],"sha256":source_sha}],
        "bronze_proofs":[{
            "table":"bronze.prg_surface_areas",
            "snapshot_id":str(bronze.current_snapshot().snapshot_id),
            "rows":rows,
            "data_files":1,
            "data_bytes":parquet_size,
            "source_sha256_column":"source_file_sha256",
        }],
        "archive":[{
            "source_key":item["key"],"archive_key":archive_key,
            "size":archive_size,"sha256":archive_sha,"compression":"native_xlsx",
        }],
        "archive_readback_verified":True,
        "landing_delete_eligible":len(blockers)==0,
        "blocking_control_references":blockers,
        "silver_table":"silver.prg_surface_areas",
        "gold_table":"gold.fact_prg_surface_areas",
        "observed_at_utc":datetime.now(timezone.utc).isoformat(),
    }
    put_immutable_json(s3,lakehouse_bucket,receipt_key,receipt)

    with bronze.transaction() as tx:
        tx.set_properties(**{"zohelo.lineage.receipt":receipt_uri})

    deleted=False
    if not blockers:
        # Revalidate Archive immediately before deleting the exact modeled XLSX.
        check_sha,check_size=hash_object(s3,landing_bucket,archive_key)
        if check_sha!=source_sha or check_size!=item["size"]:
            raise RuntimeError("PRG XLSX Archive changed before Landing deletion")
        s3.delete_object(Bucket=landing_bucket,Key=item["key"])
        if object_exists(s3,landing_bucket,item["key"]):
            raise RuntimeError("PRG XLSX Landing deletion did not complete")
        deleted=True
        lifecycle_key=f"06_control/lifecycle/gugik_prg/surface-area-{source_sha}.json"
        put_immutable_json(s3,lakehouse_bucket,lifecycle_key,{
            "format_version":1,
            "kind":"landing_archive_finalization_receipt",
            "source_id":SOURCE,
            "dataset_id":"prg_surface_areas",
            "lineage_receipt":receipt_uri,
            "landing_deleted":True,
            "deleted_landing_key":item["key"],
            "archive_key":archive_key,
            "sha256":source_sha,
            "size":item["size"],
            "completed_at_utc":datetime.now(timezone.utc).isoformat(),
        })

    print(json.dumps({
        "result":"pass",
        "source":SOURCE,
        "dataset":"prg_surface_areas",
        "source_rows":rows,
        "source_sha256":source_sha,
        "bronze_table":"bronze.prg_surface_areas",
        "silver_table":"silver.prg_surface_areas",
        "gold_table":"gold.fact_prg_surface_areas",
        "archive_verified":True,
        "landing_deleted":deleted,
        "blocking_control_references":blockers,
        "boundary_zip_untouched":True,
        "google_drive_accessed":False,
    },sort_keys=True))
    return 0


if __name__=="__main__":
    raise SystemExit(main())
